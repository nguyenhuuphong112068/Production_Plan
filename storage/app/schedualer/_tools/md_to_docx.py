"""Chuyển file Nguyên lý (markdown) sang Word (.docx) định dạng gửi khách hàng.

Dùng:  python md_to_docx.py <file.md> [<file.docx>]
Sau đó chạy docx_to_pdf.ps1 để cập nhật mục lục và xuất PDF bằng Microsoft Word.
Cần: pip install python-docx

Markdown hỗ trợ (đúng phần tài liệu Nguyên lý dùng):
  # Tiêu đề               → trang bìa
  Phiên bản X.Y · dd/mm/yyyy (dòng ngay sau tiêu đề) → bìa, footer
  đoạn văn trước "##" đầu tiên → khung giới thiệu
  ## / ### / ####          → Heading 1 / 2 / 3 ("## Phần ..." sang trang mới)
  đoạn văn có **đậm**, danh sách "- " và "1. ", bảng "| ... |"
  ![chú thích](file.png)   → hình có đánh số
  > ...                    → khung ví dụ / ghi chú
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

FONT = 'Arial'
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
INK = RGBColor(0x26, 0x26, 0x26)
GRAY = RGBColor(0x6B, 0x72, 0x80)
HEX_NAVY, HEX_BLUE, HEX_BAND, HEX_TINT, HEX_LINE = '1F3A5F', '2E75B6', 'F3F6FA', 'EAF1F8', 'C9D3DF'
TEXT_WIDTH_CM = 16.7
SUBTITLE = 'Hệ thống PMS – Lập lịch sản xuất'
AUDIENCE = 'Bộ phận Kế hoạch sản xuất'


# ---------------------------------------------------------------- tiện ích XML

def fonts(run_or_style_rpr, name=FONT):
    """Đặt font cho mọi bảng mã, bỏ font theo theme (theme ghi đè font tiêu đề)."""
    rfonts = run_or_style_rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        run_or_style_rpr.insert(0, rfonts)
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rfonts.set(qn(attr), name)


def run(paragraph, text, size=None, bold=None, italic=None, color=None):
    r = paragraph.add_run(text)
    fonts(r._element.get_or_add_rPr())
    if size:
        r.font.size = Pt(size)
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return r


def inline(paragraph, text, size=None, color=None):
    for i, part in enumerate(re.split(r'\*\*', text)):
        if part:
            run(paragraph, part, size=size, bold=(i % 2 == 1) or None, color=color)


def border(element_pr, side, size=8, color=HEX_BLUE, space=4, tag='w:pBdr'):
    bdr = element_pr.find(qn(tag))
    if bdr is None:
        bdr = OxmlElement(tag)
        element_pr.append(bdr)
    b = OxmlElement(f'w:{side}')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), str(size))
    b.set(qn('w:space'), str(space))
    b.set(qn('w:color'), color)
    bdr.append(b)


def shade(cell, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    cell._element.get_or_add_tcPr().append(shd)


def cell_borders(cell, **sides):
    """sides: top/left/bottom/right = (size, color) hoặc None để bỏ."""
    tc_pr = cell._element.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{side}')
        spec = sides.get(side)
        if spec is None:
            el.set(qn('w:val'), 'nil')
        else:
            el.set(qn('w:val'), 'single')
            el.set(qn('w:sz'), str(spec[0]))
            el.set(qn('w:color'), spec[1])
        borders.append(el)
    tc_pr.append(borders)


def table_margins(table, top=70, bottom=70, left=110, right=110):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for side, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(v))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tbl_pr.append(mar)


def field(paragraph, instr, size=None, color=None, placeholder=''):
    """Chèn field Word (PAGE, NUMPAGES, TOC)."""
    def fld(kind):
        r = paragraph.add_run()
        fonts(r._element.get_or_add_rPr())
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        el = OxmlElement('w:fldChar')
        el.set(qn('w:fldCharType'), kind)
        if kind == 'begin':
            el.set(qn('w:dirty'), 'true')
        r._element.append(el)
        return r
    fld('begin')
    r = paragraph.add_run()
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = f' {instr} '
    r._element.append(it)
    fld('separate')
    run(paragraph, placeholder, size=size, color=color)
    fld('end')


def spacing(paragraph, before=0, after=6, line=1.15):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


# ---------------------------------------------------------------- kiểu chữ

def setup_styles(doc):
    normal = doc.styles['Normal']
    fonts(normal.element.get_or_add_rPr())
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    specs = {
        'Heading 1': dict(size=17, color=NAVY, before=6, after=10, border=True),
        'Heading 2': dict(size=13, color=BLUE, before=16, after=6, border=False),
        'Heading 3': dict(size=11, color=NAVY, before=10, after=4, border=False),
    }
    for name, s in specs.items():
        st = doc.styles[name]
        fonts(st.element.get_or_add_rPr())
        st.font.size = Pt(s['size'])
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = s['color']
        pf = st.paragraph_format
        pf.space_before = Pt(s['before'])
        pf.space_after = Pt(s['after'])
        pf.keep_with_next = True
        pf.line_spacing = 1.1
        if s['border']:
            border(st.element.get_or_add_pPr(), 'bottom', size=12, color=HEX_BLUE, space=6)

    # Mục lục: TOC 1 đậm xanh, TOC 2 thụt vào
    for name, size, bold, indent in (('TOC 1', 11, True, 0), ('TOC 2', 10.5, False, 0.6)):
        try:
            st = doc.styles[name]
        except KeyError:
            st = doc.styles.add_style(name, 1)
        fonts(st.element.get_or_add_rPr())
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = NAVY if bold else INK
        st.paragraph_format.left_indent = Cm(indent)
        st.paragraph_format.space_before = Pt(8 if bold else 2)
        st.paragraph_format.space_after = Pt(2)


# ---------------------------------------------------------------- khối nội dung

def add_list_item(container, marker, text, level_indent=0.75):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(level_indent)
    pf.first_line_indent = Cm(-0.5)
    pf.tab_stops.add_tab_stop(Cm(level_indent))
    spacing(p, after=3)
    run(p, marker + '\t', color=BLUE, bold=True)
    inline(p, text)
    return p


def keep_previous_with_next(doc):
    """Câu dẫn / tiêu đề ngay trước bảng hoặc khung không bị bỏ lại cuối trang trước."""
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True


def add_table(doc, rows):
    keep_previous_with_next(doc)
    header, body = rows[0], rows[1:]
    ncol = len(header)
    # Độ rộng cột theo độ dài chữ trung bình (tối thiểu 14 %); cột đầu ngắn (mã, tên nhóm) thì in đậm
    avg = [max(8.0, sum(len(r[c]) for r in rows) / len(rows)) for c in range(ncol)]
    widths = [max(0.14, a / sum(avg)) for a in avg]
    widths = [w / sum(widths) for w in widths]
    bold_first = max(len(r[0]) for r in body) <= 28
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_margins(table)
    line = (4, HEX_LINE)
    for r, cells in enumerate([header] + body):
        tr_pr = table.rows[r]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement('w:cantSplit'))
        if r == 0:
            tr_pr.append(OxmlElement('w:tblHeader'))   # lặp hàng tiêu đề khi sang trang
        for c, text in enumerate(cells):
            cell = table.cell(r, c)
            cell.width = Cm(TEXT_WIDTH_CM * widths[c])
            cell_borders(cell, top=line, bottom=line, left=line, right=line)
            p = cell.paragraphs[0]
            spacing(p, after=0, line=1.1)
            if r == 0:
                shade(cell, HEX_NAVY)
                run(p, text, size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            else:
                if r % 2 == 0:
                    shade(cell, HEX_BAND)
                inline(p, text, size=9.5)
                if c == 0 and bold_first:
                    for rr in p.runs:
                        rr.bold = True
    after = doc.add_paragraph()
    spacing(after, after=4)


def add_box(doc, lines, fill=HEX_TINT):
    """Khung nổi bật: 1 ô, nền nhạt, viền trái xanh đậm."""
    keep_previous_with_next(doc)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_margins(table, top=120, bottom=120, left=220, right=180)
    cell = table.cell(0, 0)
    cell.width = Cm(TEXT_WIDTH_CM)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))   # không cắt khung sang 2 trang
    shade(cell, fill)
    cell_borders(cell, left=(24, HEX_BLUE))
    first = True
    for line in lines:
        if not line.strip():
            continue
        m = re.match(r'^(- |(\d+)\. )(.*)$', line)
        if m:
            add_list_item(cell, '•' if m.group(2) is None else m.group(2) + '.', m.group(3))
            continue
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        spacing(p, after=4)
        inline(p, line)
        first = False
    after = doc.add_paragraph()
    spacing(after, after=4)


# ---------------------------------------------------------------- bìa, mục lục, header/footer

def cover(doc, title, version, date):
    for _ in range(3):
        spacing(doc.add_paragraph(), after=12)
    p = doc.add_paragraph()
    spacing(p, after=10)
    run(p, 'TÀI LIỆU GIỚI THIỆU NGUYÊN LÝ', size=10.5, bold=True, color=BLUE)

    # Tách tiêu đề thành 2 dòng ở chữ "và" cho cân đối
    parts = re.split(r'\s+(?=và\s)', title, maxsplit=1)
    for i, part in enumerate(parts):
        p = doc.add_paragraph()
        spacing(p, after=0, line=1.05)
        run(p, part, size=28, bold=True, color=NAVY)
    p = doc.add_paragraph()
    spacing(p, before=14, after=24)
    border(p._element.get_or_add_pPr(), 'bottom', size=24, color=HEX_BLUE, space=10)
    run(p, SUBTITLE, size=13, color=GRAY)

    info = [('Phiên bản', version), ('Ngày phát hành', date),
            ('Phạm vi', 'Sắp lịch tự động và Kiểm soát tồn bán thành phẩm'), ('Đối tượng', AUDIENCE)]
    table = doc.add_table(rows=len(info), cols=2)
    table.autofit = False
    table_margins(table, top=60, bottom=60, left=0, right=80)
    for r, (k, v) in enumerate(info):
        a, b = table.cell(r, 0), table.cell(r, 1)
        a.width, b.width = Cm(4.2), Cm(TEXT_WIDTH_CM - 4.2)
        for c in (a, b):
            cell_borders(c, bottom=(4, HEX_LINE))
        spacing(a.paragraphs[0], after=0)
        spacing(b.paragraphs[0], after=0)
        run(a.paragraphs[0], k, size=10, bold=True, color=GRAY)
        run(b.paragraphs[0], v, size=10.5, color=INK)

    for _ in range(9):
        spacing(doc.add_paragraph(), after=12)
    p = doc.add_paragraph()
    spacing(p, after=0)
    run(p, 'Tài liệu mô tả nguyên lý hoạt động của hệ thống tại ngày phát hành. '
           'Các số liệu trong ví dụ là giả định, chỉ để minh hoạ.', size=8.5, italic=True, color=GRAY)
    p.add_run().add_break(WD_BREAK.PAGE)


def toc_page(doc):
    p = doc.add_paragraph()
    spacing(p, after=14)
    border(p._element.get_or_add_pPr(), 'bottom', size=12, color=HEX_BLUE, space=6)
    run(p, 'Mục lục', size=17, bold=True, color=NAVY)
    p = doc.add_paragraph()
    field(p, 'TOC \\o "1-2" \\h \\z \\u', placeholder='Mở bằng Word và nhấn F9 để cập nhật mục lục.')
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def header_footer(section, title, version, date):
    section.different_first_page_header_footer = True
    tab = Cm(TEXT_WIDTH_CM)

    def clear_style_tabs(paragraph):
        # Kiểu Header/Footer có sẵn tab giữa 3,25" và phải 6,5" (khổ Letter) làm chữ không căn phải
        for pos in (Inches(3.25), Inches(6.5)):
            paragraph.paragraph_format.tab_stops.add_tab_stop(pos, WD_TAB_ALIGNMENT.CLEAR)

    hp = section.header.paragraphs[0]
    clear_style_tabs(hp)
    hp.paragraph_format.tab_stops.add_tab_stop(tab, WD_TAB_ALIGNMENT.RIGHT)
    border(hp._element.get_or_add_pPr(), 'bottom', size=4, color=HEX_LINE, space=4)
    run(hp, SUBTITLE, size=8.5, color=GRAY)
    run(hp, '\t' + title, size=8.5, color=GRAY)

    fp = section.footer.paragraphs[0]
    clear_style_tabs(fp)
    fp.paragraph_format.tab_stops.add_tab_stop(tab, WD_TAB_ALIGNMENT.RIGHT)
    border(fp._element.get_or_add_pPr(), 'top', size=4, color=HEX_LINE, space=4)
    run(fp, f'Phiên bản {version} · {date}', size=8.5, color=GRAY)
    run(fp, '\tTrang ', size=8.5, color=GRAY)
    field(fp, 'PAGE', size=8.5, color=GRAY, placeholder='1')
    run(fp, ' / ', size=8.5, color=GRAY)
    field(fp, 'NUMPAGES', size=8.5, color=GRAY, placeholder='1')


# ---------------------------------------------------------------- chuyển đổi

def convert(md_path: Path, docx_path: Path):
    lines = md_path.read_text(encoding='utf-8').splitlines()
    title = next(l[2:].strip() for l in lines if l.startswith('# '))
    meta = next((l for l in lines if l.startswith('Phiên bản')), '')
    m = re.match(r'Phiên bản\s+([\d.]+)\s*·\s*(.+)', meta)
    version, date = (m.group(1), m.group(2).strip()) if m else ('', '')

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    sec.left_margin, sec.right_margin = Cm(2.3), Cm(2.0)
    sec.header_distance = sec.footer_distance = Cm(1.1)
    setup_styles(doc)
    doc.core_properties.title = title
    doc.core_properties.subject = SUBTITLE
    doc.core_properties.author = 'PMS'
    doc.core_properties.comments = f'Phiên bản {version} · {date}'

    cover(doc, title, version, date)
    toc_page(doc)
    header_footer(sec, title, version, date)

    # Đoạn giới thiệu (trước "##" đầu tiên) → khung
    first_h2 = next(i for i, l in enumerate(lines) if l.startswith('## '))
    intro = [l for l in lines[:first_h2] if l.strip() and not l.startswith('# ') and l != meta]
    figure_no = 0
    i = first_h2
    pending_intro = intro
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith('#### '):
            doc.add_paragraph(line[5:], style='Heading 3')
        elif line.startswith('### '):
            doc.add_paragraph(line[4:], style='Heading 2')
        elif line.startswith('## '):
            h = doc.add_paragraph(line[3:], style='Heading 1')
            if line.startswith('## Phần'):   # mỗi Phần bắt đầu trang mới
                h.paragraph_format.page_break_before = True
            if pending_intro:
                add_box(doc, pending_intro)
                pending_intro = None
        elif line.startswith('!['):
            mm = re.match(r'!\[(.*?)\]\((.*?)\)', line)
            figure_no += 1
            doc.add_picture(str(md_path.parent / mm.group(2)), width=Cm(16.2))
            pic = doc.paragraphs[-1]
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            spacing(pic, before=6, after=2)
            pic.paragraph_format.keep_with_next = True
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            spacing(cap, after=10)
            run(cap, f'Hình {figure_no}. ', size=9, bold=True, color=GRAY)
            run(cap, mm.group(1).replace(' · ', ' – '), size=9, italic=True, color=GRAY)
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'-{3,}', c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        elif line.startswith('>'):
            block = []
            while i < len(lines) and lines[i].startswith('>'):
                block.append(lines[i][1:].lstrip())
                i += 1
            add_box(doc, block)
            continue
        elif re.match(r'^(- |\d+\. )', line):
            mm = re.match(r'^(- |(\d+)\. )(.*)$', line)
            add_list_item(doc, '•' if mm.group(2) is None else mm.group(2) + '.', mm.group(3))
        else:
            p = doc.add_paragraph()
            inline(p, line)
        i += 1

    doc.save(str(docx_path))


if __name__ == '__main__':
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix('.docx')
    convert(src, dst)
    print(f'Đã tạo {dst}')
